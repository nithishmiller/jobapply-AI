import pytest
import sys
import os
import tempfile
from pathlib import Path
from fastapi.testclient import TestClient

# Add project root to path
project_root = Path(__file__).parent.parent
sys.path.insert(0, str(project_root))

from backend.main import app
from ai.parser import CVParser
from docx import Document

parser = CVParser()

def test_root_endpoint():
    client = TestClient(app)
    response = client.get("/")
    assert response.status_code == 200
    data = response.json()
    assert data["app"] == "JobApply AI"
    assert data["status"] == "online"

def test_health_endpoint():
    client = TestClient(app)
    response = client.get("/health")
    assert response.status_code == 200
    data = response.json()
    assert data["status"] == "healthy"

def test_upload_valid_pdf():
    client = TestClient(app)
    # Generate a realistic PDF using parser helper
    pdf_path = "test_cv.pdf"
    pdf_content = """
John Doe
Email: john@example.com
Phone: 555-123-4567

Technical Stack:
Python, Django, MySQL, SQLite, HTML, CSS, JavaScript, REST APIs, Git, GitHub

Education:
Bachelor of Science in Computer Science
University of Example (2016-2020)
"""
    parser.create_test_pdf(pdf_path, pdf_content)
    with open(pdf_path, "rb") as f:
        response = client.post(
            "/upload-cv",
            files={"file": ("test_cv.pdf", f, "application/pdf")}
        )
    assert response.status_code == 200
    data = response.json()
    assert data["status"] == "uploaded"
    assert data["filename"] == "test_cv.pdf"
    assert data["parsed"] is True
    os.remove(pdf_path)

def test_upload_valid_docx():
    client = TestClient(app)
    # Generate a realistic DOCX programmatically
    docx_path = "test_cv.docx"
    doc = Document()
    doc.add_heading("John Doe", 0)
    doc.add_paragraph("Email: john@example.com")
    doc.add_paragraph("Phone: 555-123-4567")
    doc.add_heading("Technical Stack", level=1)
    doc.add_paragraph("Python, Django, MySQL, SQLite, HTML, CSS, JavaScript, REST APIs, Git, GitHub")
    doc.add_heading("Education", level=1)
    doc.add_paragraph("Bachelor of Science in Computer Science")
    doc.add_paragraph("University of Example (2016-2020)")
    doc.save(docx_path)
    with open(docx_path, "rb") as f:
        response = client.post(
            "/upload-cv",
            files={"file": ("test_cv.docx", f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        )
    assert response.status_code == 200
    data = response.json()
    assert data["status"] == "uploaded"
    assert data["filename"] == "test_cv.docx"
    assert data["parsed"] is True
    os.remove(docx_path)

def test_upload_invalid_file():
    client = TestClient(app)
    fd, path = tempfile.mkstemp(suffix=".txt")
    try:
        with os.fdopen(fd, 'wb') as tmp:
            tmp.write(b"This is not a valid CV file")
        with open(path, 'rb') as f:
            response = client.post(
                "/upload-cv",
                files={"file": ("test.txt", f, "text/plain")}
            )
        assert response.status_code == 400
        data = response.json()
        assert "error" in data
    finally:
        if os.path.exists(path):
            os.unlink(path)

def test_get_cvs_parsed_entries():
    client = TestClient(app)
    # Ensure parsed CV entries are present and have extracted data
    response = client.get("/cvs/")
    assert response.status_code == 200
    data = response.json()
    assert data["status"] == "success"
    parsed_entries = [cv for cv in data["data"] if cv.get("parsed")]
    # We expect at least the two uploads from previous tests
    assert len(parsed_entries) >= 2
    for cv in parsed_entries:
        detail = client.get(f"/cvs/{cv['id']}")
        assert detail.status_code == 200
        d = detail.json()["data"]
        assert d["contact_info"] is not None
        # Education is optional - don't require it
        # Experience is optional - don't require it
        assert d["skills"]